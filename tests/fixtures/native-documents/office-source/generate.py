"""Author synthetic input; no external document, personal data or network.

Run with python-docx, python-pptx and openpyxl installed. The archive rewrite
fixes ZIP metadata; package library versions are recorded in README.md.
"""
from datetime import datetime, timezone
from pathlib import Path
import tempfile
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.worksheet.page import PageMargins
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt

DEST = Path(__file__).resolve().parent.parent
EPOCH = datetime(2020, 1, 1, tzinfo=timezone.utc)


def stable_archive(source, target):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for name in sorted(original.namelist()):
            info = zipfile.ZipInfo(name, (2020, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            output.writestr(info, original.read(name))


with tempfile.TemporaryDirectory(prefix="baro-author-office-") as temporary:
    root = Path(temporary)
    document = Document()
    document.core_properties.author = "BARO synthetic fixture"
    document.core_properties.created = document.core_properties.modified = EPOCH
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Noto Sans CJK KR", Pt(12)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    style.element.get_or_add_rPr().append(fonts)
    document.add_paragraph("합성 자료 첫째 쪽 BLUEBOX17")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for cell, text in zip((cell for row in table.rows for cell in row.cells),
                          ["합성 표", "ORANGECIRCLE42", "자료", "확인"]):
        cell.text = text
    document.add_page_break()
    document.add_paragraph("합성 자료 둘째 쪽 END_NATIVE_DOCUMENT")
    document.save(root / "input.docx")
    stable_archive(root / "input.docx", DEST / "office-two-pages.docx")

    workbook = Workbook()
    workbook.properties.creator = "BARO synthetic fixture"
    workbook.properties.created = workbook.properties.modified = EPOCH
    for index in range(2):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = "Synthetic" + str(index + 1)
        sheet.append(["합성 자료 첫째 쪽" if index == 0 else "합성 자료 둘째 쪽",
                      "BLUEBOX17" if index == 0 else "END_NATIVE_DOCUMENT"])
        sheet.append(["합성 표", "ORANGECIRCLE42"])
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 34
        sheet.print_area = "A1:B2"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
        sheet.page_margins = PageMargins(left=.2, right=.2, top=.2, bottom=.2, header=0, footer=0)
    workbook.save(root / "input.xlsx")
    stable_archive(root / "input.xlsx", DEST / "office-two-sheets.xlsx")

    presentation = Presentation()
    presentation.core_properties.author = "BARO synthetic fixture"
    presentation.core_properties.created = presentation.core_properties.modified = EPOCH
    for index in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(9), Inches(1)).text_frame
        run = frame.paragraphs[0].add_run()
        run.text = "합성 자료 첫째 쪽 BLUEBOX17" if index == 0 else "합성 자료 둘째 쪽 END_NATIVE_DOCUMENT"
        run.font.name, run.font.size = "Noto Sans CJK KR", PptPt(20)
        table = slide.shapes.add_table(2, 2, Inches(.5), Inches(2), Inches(8), Inches(1)).table
        for cell, text in zip((table.cell(row, col) for row in range(2) for col in range(2)),
                              ["합성 표", "ORANGECIRCLE42", "자료", "확인"]):
            cell.text = text
    presentation.save(root / "input.pptx")
    stable_archive(root / "input.pptx", DEST / "office-two-slides.pptx")

(DEST / "utf8-two-pages.txt").write_bytes(
    ("\ufeff합성 자료 첫째 쪽 BLUEBOX17\f합성 자료 둘째 쪽 END_NATIVE_DOCUMENT").encode("utf-8"))
